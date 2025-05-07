import os
import re
import xlrd
import openpyxl
from datetime import datetime
from openpyxl.cell.cell import Cell
import pandas as pd

# Main folder path
main_folder_path = r"D:\RIghtstroken\Pricing\PricingExtraction\Pricing_Sheets_24-Feb-25"
output_file = "final_summary_all_folders.xlsx"

# Field patterns
patterns = {
    "Part# / Model name": r"(part\s*#|model\s*name)",
    "OPP#": r"opp\s*#?",
    "CUSTOMER": r"(customer|client)(\s*name)?",
    "Assembly cost / PPD": r"\b(assembly cost|ppd)\b",
    "Estimated BOM cost": r"\b(estimated bom cost|bom cost per unit)\b",
    "Design & Development cost": r"design and development cost",
    "Recommended Price": r"recommended price",
    "Comments to Steven": r"(comments for steven\.s|comments to steven)",
    "CREATED ON": r"created\s*on\s*[:\-]?"
}

# Clean and extract value
def clean_value(value, key, cell=None):
    if isinstance(cell, Cell) and cell.is_date:
        return cell.value.strftime("%m/%d/%Y")

    if key == "CREATED ON":
        if isinstance(value, float):
            try:
                date_value = datetime(*xlrd.xldate_as_tuple(value, 0))
                return date_value.strftime("%m/%d/%Y")
            except:
                pass
        if isinstance(value, str):
            match = re.search(r"(\d{1,2})[-/\s](\w{3,})[-/\s](\d{2,4})", value, re.IGNORECASE)
            if match:
                for fmt in ["%d-%b-%y", "%d-%b-%Y", "%d/%b/%y", "%d/%b/%Y"]:
                    try:
                        return datetime.strptime(match.group(0), fmt).strftime("%m/%d/%Y")
                    except:
                        continue
    return str(value).strip()

# Extract from .xls files
def extract_from_xls(sheet):
    extracted = {}
    for row_idx in range(sheet.nrows):
        for col_idx in range(sheet.ncols):
            cell_text = str(sheet.cell_value(row_idx, col_idx)).strip().lower()
            for key, pattern in patterns.items():
                if key not in extracted and re.search(pattern, cell_text, re.IGNORECASE):
                    cleaned = clean_value(sheet.cell_value(row_idx, col_idx), key)
                    if cleaned and cleaned.lower() != cell_text:
                        extracted[key] = cleaned
                        continue

                    next_val = ""
                    try:
                        if col_idx + 1 < sheet.ncols:
                            next_val = sheet.cell_value(row_idx, col_idx + 1)
                        elif row_idx + 1 < sheet.nrows:
                            next_val = sheet.cell_value(row_idx + 1, col_idx)
                    except:
                        pass

                    extracted[key] = clean_value(next_val, key)
    return extracted

# Extract from .xlsx files
def extract_from_xlsx(sheet):
    extracted = {}
    max_row, max_col = sheet.max_row, sheet.max_column

    for row in sheet.iter_rows():
        for cell in row:
            if cell.value:
                cell_text = str(cell.value).strip().lower()
                for key, pattern in patterns.items():
                    if key not in extracted or not extracted[key]:
                        if re.search(pattern, cell_text, re.IGNORECASE):
                            print(f"🔍 Found '{key}' in cell: {cell_text}")
                            cleaned = clean_value(cell.value, key, cell)
                            if cleaned and cleaned.lower() != cell_text:
                                extracted[key] = cleaned
                                continue

                            next_val = None
                            try:
                                if cell.column + 1 <= max_col:
                                    next_val = sheet.cell(cell.row, cell.column + 1).value
                            except:
                                pass

                            cleaned_next = clean_value(next_val, key, cell)
                            if cleaned_next:
                                extracted[key] = cleaned_next
                                if key == "CUSTOMER":
                                    print(f"✅ CUSTOMER value (next cell): {cleaned_next}")
                            elif key == "CUSTOMER":
                                try:
                                    further_next = sheet.cell(cell.row, cell.column + 2).value
                                    cleaned_further = clean_value(further_next, key, cell)
                                    if cleaned_further:
                                        extracted[key] = cleaned_further
                                        print(f"✅ CUSTOMER value (2nd next cell): {cleaned_further}")
                                except:
                                    pass
    return extracted

# Write to final Excel
with pd.ExcelWriter(output_file, engine='xlsxwriter') as writer:
    for subfolder in os.listdir(main_folder_path):
        subfolder_path = os.path.join(main_folder_path, subfolder)
        if os.path.isdir(subfolder_path):
            data = []
            print(f"\n📁 Processing folder: {subfolder}")
            for filename in os.listdir(subfolder_path):
                if filename.endswith((".xls", ".xlsx")):
                    file_path = os.path.join(subfolder_path, filename)
                    row_data = {"File Name": filename}
                    try:
                        print(f"📄 Reading file: {filename}")
                        if filename.endswith(".xls"):
                            book = xlrd.open_workbook(file_path)
                            sheet = book.sheet_by_index(0)
                            extracted = extract_from_xls(sheet)
                        else:
                            wb = openpyxl.load_workbook(file_path, data_only=True)
                            sheet = wb.active
                            extracted = extract_from_xlsx(sheet)

                        row_data.update(extracted)
                        data.append(row_data)

                        if "CUSTOMER" in extracted:
                            print(f"📝 CUSTOMER for {filename}: {extracted.get('CUSTOMER')}")
                        else:
                            print(f"⚠️ CUSTOMER NOT FOUND in {filename}")

                    except Exception as e:
                        print(f"❌ Error reading {filename} in {subfolder}: {e}")

            if data:
                df = pd.DataFrame(data)
                sheet_name = subfolder[:31]
                df.to_excel(writer, index=False, sheet_name=sheet_name)
                worksheet = writer.sheets[sheet_name]
                workbook = writer.book

                # Set column width
                for col_num in range(len(df.columns)):
                    worksheet.set_column(col_num, col_num, 25)

                # Highlight blank cells
                format_blank = workbook.add_format({'bg_color': '#FFC7CE'})
                worksheet.conditional_format(1, 0, len(df), len(df.columns) - 1, {
                    'type': 'blanks',
                    'format': format_blank
                })

print(f"\n✅ All summaries saved to {output_file}")
